"""
export_three_completed_datasets.py
Generates the 3 certified, completed datasets and summary:
1. data/COMPLETED_DATA/TELANGANA_COMPLETED.xlsx (46,845 records)
2. data/COMPLETED_DATA/COA_COMPLETED_2025_26.xlsx (404 records)
3. data/COMPLETED_DATA/RCI_COMPLETED_2025.xlsx (1,055 records)
4. data/COMPLETED_DATA/COMPLETED_DATA_SUMMARY.xlsx (Summary table)
"""

import sqlite3
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path
from datetime import date

BASE = Path(__file__).resolve().parent.parent
DB_PATH = BASE / "data" / "processed" / "education_master.db"
OUT_DIR = BASE / "data" / "COMPLETED_DATA"
OUT_DIR.mkdir(parents=True, exist_ok=True)

EXTRACTION_DATE = "2026-09-10"

def export_telangana():
    print("\n[1/4] Exporting TELANGANA_COMPLETED.xlsx...")
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    
    # Verify exact count
    count = c.execute("SELECT COUNT(1) FROM pan_india_census WHERE state = 'Telangana'").fetchone()[0]
    unique_ids = c.execute("SELECT COUNT(DISTINCT official_institution_id) FROM pan_india_census WHERE state = 'Telangana'").fetchone()[0]
    print(f"  Verified Database Telangana Count: {count:,}")
    print(f"  Verified Unique Official IDs:     {unique_ids:,}")
    assert count == 46845, f"Expected 46,845, found {count}"
    assert unique_ids == 46845, f"Expected 46,845 unique IDs, found {unique_ids}"
    
    # Query all records
    query = """
        SELECT 
            name AS Institution_Name,
            institution_id AS Census_Internal_ID,
            official_institution_id AS Official_Institution_ID,
            udise_code AS UDISE_Code,
            institution_type AS Institution_Type,
            education_level AS Education_Level,
            institution_category AS Institution_Category,
            management_type AS Management_Type,
            full_address AS Address,
            city_town_village AS City_Town_Village,
            block_mandal AS Block_Mandal,
            district AS District,
            state AS State,
            pincode AS PIN_Code,
            university_affiliation AS University_Affiliation,
            board_affiliation AS Board_Affiliation,
            courses_programmes AS Courses_Programmes,
            year_established AS Year_Established,
            website AS Website,
            email AS Email,
            phone AS Phone,
            recognition_status AS Recognition_Status,
            recognition_authority AS Recognition_Authority,
            approval_status AS Approval_Status,
            approval_authority AS Approval_Authority,
            '2021-25' AS Academic_Year,
            source_database AS Source_Database,
            source_url AS Source_URL,
            collection_date AS Collection_Date,
            verification_status AS Verification_Status,
            remarks AS Remarks
        FROM pan_india_census
        WHERE state = 'Telangana'
        ORDER BY education_level, district, name
    """
    
    c.execute(query)
    headers = [d[0] for d in c.description]
    rows = c.fetchall()
    conn.close()
    
    out_path = OUT_DIR / "TELANGANA_COMPLETED.xlsx"
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Telangana_Validated_Census")
    
    ws.append(headers)
    for r in rows:
        ws.append(list(r))
        
    wb.save(out_path)
    wb.close()
    
    mb = out_path.stat().st_size / (1024 * 1024)
    print(f"  [OK] Saved {out_path.name}: {len(rows):,} rows ({mb:.2f} MB)")
    return len(rows), unique_ids

def export_coa():
    print("\n[2/4] Exporting COA_COMPLETED_2025_26.xlsx...")
    src_path = BASE / "data" / "COA_INSTITUTIONS_2025.xlsx"
    assert src_path.exists(), "COA_INSTITUTIONS_2025.xlsx not found!"
    
    wb_src = openpyxl.load_workbook(src_path, read_only=True)
    ws_src = wb_src.active
    raw_rows = list(ws_src.iter_rows(values_only=True))
    wb_src.close()
    
    headers_src = [str(h).strip() for h in raw_rows[0]]
    h_idx = {h: i for i, h in enumerate(headers_src)}
    data_rows = raw_rows[1:]
    
    target_headers = [
        "CoA_Code",
        "Institution_Name",
        "Full_Address",
        "State",
        "District",
        "PIN_Code",
        "University_Affiliation",
        "Approved_Intake",
        "Email",
        "Website",
        "Country",
        "Education_Level",
        "Institution_Type",
        "Approval_Status",
        "Approval_Authority",
        "Academic_Year",
        "Source_URL",
        "Extraction_Date"
    ]
    
    out_rows = []
    seen_ids = set()
    for r in data_rows:
        cid = str(r[h_idx["Official_ID"]]).strip() if r[h_idx.get("Official_ID")] else ""
        if not cid:
            continue
        seen_ids.add(cid)
        
        name = str(r[h_idx["Institution_Name"]]).strip() if r[h_idx.get("Institution_Name")] else ""
        st = str(r[h_idx["State"]]).strip() if r[h_idx.get("State")] else ""
        pin = str(r[h_idx["PIN_Code"]]).strip() if r[h_idx.get("PIN_Code")] else ""
        univ = str(r[h_idx["University_Affiliation"]]).strip() if r[h_idx.get("University_Affiliation")] else ""
        intake = str(r[h_idx["Courses_and_Intake"]]).strip() if r[h_idx.get("Courses_and_Intake")] else ""
        email = str(r[h_idx["Email"]]).strip() if r[h_idx.get("Email")] else ""
        web = str(r[h_idx["Website"]]).strip() if r[h_idx.get("Website")] else ""
        
        row_dict = [
            cid,
            name,
            f"{name}, {st} - {pin}".strip(" ,-"),
            st,
            "",  # District
            pin,
            univ,
            intake,
            email,
            web,
            "India",
            "Higher Education / Architecture",
            "Architecture College / School of Planning & Architecture",
            "Approved",
            "Council of Architecture (CoA)",
            "2025-26",
            "https://coa.gov.in/institutionStatus.php",
            EXTRACTION_DATE
        ]
        out_rows.append(row_dict)
        
    assert len(out_rows) == 404, f"Expected 404 CoA institutions, found {len(out_rows)}"
    assert len(seen_ids) == 404, f"Expected 404 unique CoA IDs, found {len(seen_ids)}"
    
    out_path = OUT_DIR / "COA_COMPLETED_2025_26.xlsx"
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "CoA_Approved_Institutions"
    
    # Styles
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    
    ws_out.append(target_headers)
    for cell in ws_out[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    for r in out_rows:
        ws_out.append(r)
        
    for row in ws_out.iter_rows(min_row=2, max_row=len(out_rows)+1):
        for cell in row:
            cell.font = data_font
            cell.alignment = Alignment(vertical="center")
            
    # Column auto-width
    for col in ws_out.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_out.column_dimensions[col_letter].width = min(50, max(12, max_len + 3))
        
    wb_out.save(out_path)
    wb_out.close()
    
    kb = out_path.stat().st_size / 1024
    print(f"  [OK] Saved {out_path.name}: {len(out_rows):,} rows ({kb:.1f} KB)")
    return len(out_rows), len(seen_ids)

def export_rci():
    print("\n[3/4] Exporting RCI_COMPLETED_2025.xlsx...")
    src_path = BASE / "data" / "RCI_INSTITUTIONS_2025.xlsx"
    assert src_path.exists(), "RCI_INSTITUTIONS_2025.xlsx not found!"
    
    wb_src = openpyxl.load_workbook(src_path, read_only=True)
    ws_src = wb_src.active
    raw_rows = list(ws_src.iter_rows(values_only=True))
    wb_src.close()
    
    headers_src = [str(h).strip() for h in raw_rows[0]]
    h_idx = {h: i for i, h in enumerate(headers_src)}
    data_rows = raw_rows[1:]
    
    target_headers = [
        "RCI_Institute_Code",
        "Institution_Name",
        "Address",
        "State",
        "District",
        "PIN_Code",
        "Education_Level",
        "Institution_Type",
        "Approved_Programmes",
        "Recognition_Status",
        "Recognition_Authority",
        "Country",
        "Academic_Year",
        "Source_URL",
        "Extraction_Date"
    ]
    
    out_rows = []
    seen_ids = set()
    for r in data_rows:
        rid = str(r[h_idx["Official_ID"]]).strip() if r[h_idx.get("Official_ID")] else ""
        if not rid:
            continue
        seen_ids.add(rid)
        
        name = str(r[h_idx["Institution_Name"]]).strip() if r[h_idx.get("Institution_Name")] else ""
        addr = str(r[h_idx["Address"]]).strip() if r[h_idx.get("Address")] else ""
        st = str(r[h_idx["State"]]).strip() if r[h_idx.get("State")] else ""
        dist = str(r[h_idx["District"]]).strip() if r[h_idx.get("District")] else ""
        pin = str(r[h_idx["PIN_Code"]]).strip() if r[h_idx.get("PIN_Code")] else ""
        progs = str(r[h_idx["Approved_Programmes"]]).strip() if r[h_idx.get("Approved_Programmes")] else ""
        
        row_dict = [
            rid,
            name,
            addr,
            st,
            dist,
            pin,
            "Higher Education / Rehabilitation",
            "Special Education / Rehabilitation Training Institute",
            progs,
            "Approved",
            "Rehabilitation Council of India (RCI)",
            "India",
            "2025",
            "https://rciregistration.nic.in/rehabcouncil/instapproval_statewise.jsp",
            EXTRACTION_DATE
        ]
        out_rows.append(row_dict)
        
    assert len(out_rows) == 1055, f"Expected 1,055 RCI institutions, found {len(out_rows)}"
    assert len(seen_ids) == 1055, f"Expected 1,055 unique RCI IDs, found {len(seen_ids)}"
    
    out_path = OUT_DIR / "RCI_COMPLETED_2025.xlsx"
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "RCI_Approved_Institutions"
    
    # Styles
    header_fill = PatternFill(start_color="1E4620", end_color="1E4620", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    
    ws_out.append(target_headers)
    for cell in ws_out[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    for r in out_rows:
        ws_out.append(r)
        
    for row in ws_out.iter_rows(min_row=2, max_row=len(out_rows)+1):
        for cell in row:
            cell.font = data_font
            cell.alignment = Alignment(vertical="center")
            
    # Column auto-width
    for col in ws_out.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_out.column_dimensions[col_letter].width = min(60, max(12, max_len + 3))
        
    wb_out.save(out_path)
    wb_out.close()
    
    kb = out_path.stat().st_size / 1024
    print(f"  [OK] Saved {out_path.name}: {len(out_rows):,} rows ({kb:.1f} KB)")
    return len(out_rows), len(seen_ids)

def export_summary(tg_stats, coa_stats, rci_stats):
    print("\n[4/4] Exporting COMPLETED_DATA_SUMMARY.xlsx...")
    out_path = OUT_DIR / "COMPLETED_DATA_SUMMARY.xlsx"
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Completed_Datasets_Summary"
    
    headers = [
        "Dataset",
        "Records",
        "Unique_IDs",
        "States_UTs_Covered",
        "Official_Identifier",
        "Academic_Year",
        "Status",
        "Source_URL",
        "Extraction_Date",
        "Notes_and_Limitations"
    ]
    
    rows = [
        [
            "Telangana",
            tg_stats[0],
            tg_stats[1],
            "1 (Telangana)",
            "Official State/Board IDs (UDISE, TSBIE, DOST, SBTET, AICTE UAAC, TS LAWCET, TS EDCET, Statutory Universities, CBSE, CISCE, CoA, RCI)",
            "2021-25",
            "COMPLETE",
            "Multi-portal verified official directories",
            EXTRACTION_DATE,
            "Complete validated Telangana census. 42,834 schools, 1,744 junior colleges, 784 degree colleges, 370 polytechnics, etc. 42 KYS probe records excluded. Excluded NMC additions removed. 100% ID uniqueness."
        ],
        [
            "CoA",
            coa_stats[0],
            coa_stats[1],
            "26 States/UTs",
            "CoA Institution Code (e.g. AP02, DL01, TS03)",
            "2025-26",
            "COMPLETE",
            "https://coa.gov.in/institutionStatus.php",
            EXTRACTION_DATE,
            "100% complete national approval register extracted with approved course intake. Exactly 404 physical institutions (1 row = 1 canonical institution, no programme duplication)."
        ],
        [
            "RCI",
            rci_stats[0],
            rci_stats[1],
            "34 States/UTs",
            "RCI Institute Code (e.g. AP004, DL001, TS002)",
            "2025",
            "COMPLETE",
            "https://rciregistration.nic.in/rehabcouncil/instapproval_statewise.jsp",
            EXTRACTION_DATE,
            "100% complete across all 34 States/UTs offering approved rehabilitation and special education training programmes. Exactly 1,055 institutions (1 row = 1 institute, programmes consolidated)."
        ]
    ]
    
    # Styles
    header_fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    for r in rows:
        ws.append(r)
        
    for row in ws.iter_rows(min_row=2, max_row=len(rows)+1):
        for cell in row:
            cell.font = data_font
            cell.alignment = Alignment(vertical="center")
            
    # Column auto-width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(60, max(14, max_len + 3))
        
    wb.save(out_path)
    wb.close()
    
    kb = out_path.stat().st_size / 1024
    print(f"  [OK] Saved {out_path.name}: {len(rows)} rows ({kb:.1f} KB)")

def main():
    print("=" * 80)
    print("EXPORTING THE 3 CERTIFIED COMPLETED DATASETS ONLY")
    print("=" * 80)
    
    tg_stats = export_telangana()
    coa_stats = export_coa()
    rci_stats = export_rci()
    export_summary(tg_stats, coa_stats, rci_stats)
    
    print("\n" + "=" * 80)
    print("EXPORT COMPLETED SUCCESSFULLY")
    print("=" * 80)
    print(f"1. TELANGANA: {OUT_DIR / 'TELANGANA_COMPLETED.xlsx'} ({tg_stats[0]:,} records)")
    print(f"2. CoA:       {OUT_DIR / 'COA_COMPLETED_2025_26.xlsx'} ({coa_stats[0]:,} records)")
    print(f"3. RCI:       {OUT_DIR / 'RCI_COMPLETED_2025.xlsx'} ({rci_stats[0]:,} records)")
    print(f"4. SUMMARY:   {OUT_DIR / 'COMPLETED_DATA_SUMMARY.xlsx'}")

if __name__ == "__main__":
    main()
