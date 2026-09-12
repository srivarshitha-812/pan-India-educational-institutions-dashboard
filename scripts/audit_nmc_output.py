import openpyxl
from collections import Counter

def audit_file(filepath):
    print(f"Auditing {filepath} ...")
    wb = openpyxl.load_workbook(filepath, read_only=True)
    sheetname = wb.sheetnames[0]
    ws = wb[sheetname]
    
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    data = rows[1:]
    
    print(f"Sheet Name: {sheetname}")
    print(f"Header: {header}")
    print(f"Total Rows: {len(data)}")
    
    col_idx = {name: i for i, name in enumerate(header)}
    
    # Check ID uniqueness
    id_col = col_idx.get('NMC_College_ID')
    if id_col is not None:
        ids = [r[id_col] for r in data]
        unique_ids = set(ids)
        print(f"Unique NMC_College_ID: {len(unique_ids)}")
        if len(ids) == len(unique_ids):
            print("100% Unique - 0 Duplicates!")
        else:
            print(f"Duplicates found: {len(ids) - len(unique_ids)}")
            
    # If course file, check seat sum and recognition
    if 'Annual_Intake' in col_idx:
        seat_col = col_idx['Annual_Intake']
        total_seats = sum(int(r[seat_col] or 0) for r in data)
        print(f"Total Annual Intake Seats: {total_seats:,}")
        
    if 'Recognition_Status' in col_idx:
        rec_col = col_idx['Recognition_Status']
        rec_counts = Counter(r[rec_col] for r in data)
        print("Recognition Status counts:", dict(rec_counts))
        
    # Print first 2 data rows
    print("\nSample row 1:")
    for col_name, idx in col_idx.items():
        print(f"  {col_name}: {repr(data[0][idx])}")
    print("\nSample row 2:")
    for col_name, idx in col_idx.items():
        print(f"  {col_name}: {repr(data[1][idx])}")
    print("\n" + "="*70 + "\n")

audit_file('data/NMC_COLLEGES_2026_27.xlsx')
audit_file('data/NMC_COURSES_2026_27.xlsx')
