import json, re
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

print("Loading data...")
with open('data/nmc_all_colleges_api.json', encoding='utf-8') as f:
    colleges_raw = json.load(f)

with open('data/nmc_all_courses_api.json', encoding='utf-8') as f:
    courses_raw = json.load(f)

print(f"Loaded {len(colleges_raw)} colleges from API and {len(courses_raw)} courses from API.")

# Map collegeId -> college object from API
col_api_map = {c['collegeId']: c for c in colleges_raw}

# Helper to extract NMC_College_ID and clean Institution_Name
def parse_college_id_and_name(cid, raw_name, col_obj):
    raw_name = raw_name.strip()
    m = re.match(r'^([A-Z0-9/_-]+):\s*(.*)$', raw_name)
    if m:
        code = m.group(1).strip()
        name = m.group(2).strip()
    else:
        # For institutes without prefix (like AIIMS), use official collegeCode from NMC database
        code = col_obj.get('collegeCode') if col_obj and col_obj.get('collegeCode') else f"NMC-{cid:04d}"
        name = raw_name
    return code, name

# Helper to clean address
def clean_address(col_obj):
    if not col_obj:
        return ""
    addr = col_obj.get('address') or ""
    # replace newlines with comma and space
    addr = re.sub(r'[\r\n]+', ', ', addr).strip()
    addr = re.sub(r'\s+', ' ', addr)
    addr = re.sub(r',\s*,', ',', addr).strip(' ,')
    
    city = (col_obj.get('city') or "").strip()
    pincode = (col_obj.get('pincode') or "").strip()
    
    if city and city.lower() not in addr.lower():
        addr = f"{addr}, {city}" if addr else city
    if pincode and pincode not in addr:
        addr = f"{addr} - {pincode}" if addr else pincode
        
    return addr.strip(' ,')

# Helper to clean state name
def clean_state(s):
    if not s:
        return ""
    s = s.strip()
    # Replace double spaces
    s = re.sub(r'\s+', ' ', s)
    # Fix known NMC typos like "Andaman  Nicobar Islands" -> "Andaman & Nicobar Islands"
    if 'Andaman' in s and 'Nicobar' in s:
        return "Andaman & Nicobar Islands"
    if s == "Chattisgarh":
        return "Chhattisgarh"
    if s == "Orissa":
        return "Odisha"
    if s == "Pondicherry":
        return "Puducherry"
    return s

# Helper to clean university name
def clean_univ(u):
    if not u:
        return ""
    u = u.strip()
    u = re.sub(r'\s+', ' ', u)
    return u

# Helper to map recognition status
def map_recognition(rec):
    rec = (rec or "").strip().upper()
    if rec == 'R':
        return 'Recognized'
    elif rec == 'P':
        return 'Permitted'
    elif rec == 'O':
        return 'Others / Granted'
    return rec if rec else 'N/A'

# ─────────────────────────────────────────────────────────────────────────────
# 1. BUILD NMC_COLLEGES_2026_27
# ─────────────────────────────────────────────────────────────────────────────
print("\nBuilding Colleges Dataset...")

# Aggregate courses info by collegeId
college_courses_map = {}
for c in courses_raw:
    cid = c['collegeId']
    college_courses_map.setdefault(cid, []).append(c)

unique_college_ids = sorted(college_courses_map.keys())
print(f"Unique colleges offering courses: {len(unique_college_ids)}")

college_rows = []
for cid in unique_college_ids:
    c_list = college_courses_map[cid]
    first_course = c_list[0]
    col_obj = col_api_map.get(cid, {})
    
    raw_col_name = first_course.get('collegeName', '')
    code, name = parse_college_id_and_name(cid, raw_col_name, col_obj)
    
    # State
    state = clean_state(first_course.get('stateName', ''))
    
    # University: collect all distinct universities affiliated to this college
    all_univs = sorted(set(clean_univ(x.get('univName', '')) for x in c_list if x.get('univName')))
    univ_str = "; ".join(all_univs) if all_univs else clean_univ(col_obj.get('university1') or "")
    
    # Management
    mgmt = first_course.get('managementText') or col_obj.get('managementText') or ""
    mgmt = mgmt.strip()
    if mgmt.lower() == 'govt.':
        mgmt = 'Govt.'
    elif mgmt.lower() == 'private':
        mgmt = 'Private'
    elif mgmt.lower() == 'government':
        mgmt = 'Govt.'
        
    # Year of Inception
    yoi = col_obj.get('yearOfInc') or first_course.get('year') or ""
    yoi = str(yoi).strip()
    
    # Address
    address = clean_address(col_obj)
    
    # Status
    status = "Active / LIVE" if col_obj.get('onlineFlag') == 'LIVE' else "Active"
    
    college_rows.append({
        'NMC_College_ID': code,
        'Institution_Name': name,
        'Address': address,
        'State': state,
        'University': univ_str,
        'Management': mgmt,
        'Year_of_Inception': yoi,
        'Status': status,
        'Source_URL': 'https://www.nmc.org.in/information-desk/college-and-course-search/',
        'Academic_Year': '2026-27'
    })

df_colleges = pd.DataFrame(college_rows)
print(f"Total College Rows: {len(df_colleges)}")
print(f"Unique NMC_College_ID count: {df_colleges['NMC_College_ID'].nunique()}")
assert len(df_colleges) == df_colleges['NMC_College_ID'].nunique(), "Collision in NMC_College_ID!"

# ─────────────────────────────────────────────────────────────────────────────
# 2. BUILD NMC_COURSES_2026_27
# ─────────────────────────────────────────────────────────────────────────────
print("\nBuilding Courses Dataset...")

# Build mapping from cid to code and clean name
cid_to_meta = {}
for cid in unique_college_ids:
    first_course = college_courses_map[cid][0]
    col_obj = col_api_map.get(cid, {})
    code, name = parse_college_id_and_name(cid, first_course.get('collegeName', ''), col_obj)
    cid_to_meta[cid] = (code, name)

course_rows = []
for c in courses_raw:
    cid = c['collegeId']
    code, name = cid_to_meta[cid]
    
    cname = c.get('courseName', '').strip()
    state = clean_state(c.get('stateName', ''))
    univ = clean_univ(c.get('univName', ''))
    
    mgmt = c.get('managementText', '').strip()
    if mgmt.lower() == 'govt.':
        mgmt = 'Govt.'
    elif mgmt.lower() == 'private':
        mgmt = 'Private'
    elif mgmt.lower() == 'government':
        mgmt = 'Govt.'
        
    seat = int(c.get('seat', 0) or 0)
    rec_status = map_recognition(c.get('recognization'))
    
    course_rows.append({
        'NMC_College_ID': code,
        'Institution_Name': name,
        'Course_Name': cname,
        'State': state,
        'University': univ,
        'Management': mgmt,
        'Annual_Intake': seat,
        'Recognition_Status': rec_status,
        'Academic_Year': '2026-27'
    })

df_courses = pd.DataFrame(course_rows)
print(f"Total Course Rows: {len(df_courses)}")
print(f"Total Annual Intake Seats: {df_courses['Annual_Intake'].sum():,}")

# ─────────────────────────────────────────────────────────────────────────────
# 3. WRITE EXCEL FILES WITH PROFESSIONAL FORMATTING
# ─────────────────────────────────────────────────────────────────────────────
def save_styled_excel(df, filepath, sheet_name="Sheet1"):
    print(f"Writing {filepath}...")
    with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name=sheet_name, index=False)
        
    # Apply styling
    wb = openpyxl.load_workbook(filepath)
    ws = wb[sheet_name]
    
    # Header styling
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    border_thin = Border(
        left=Side(style='thin', color="D9D9D9"),
        right=Side(style='thin', color="D9D9D9"),
        top=Side(style='thin', color="D9D9D9"),
        bottom=Side(style='thin', color="D9D9D9")
    )
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    ws.row_dimensions[1].height = 28
    
    # Data rows styling
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.font = data_font
            cell.border = border_thin
            # Numbers alignment
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
    
    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col[:100])
        header_len = len(str(col[0].value or ''))
        width = max(max_len + 3, header_len + 3)
        ws.column_dimensions[col_letter].width = min(width, 50)
        
    wb.save(filepath)
    print(f"Successfully saved and formatted {filepath}")

save_styled_excel(df_colleges, 'data/NMC_COLLEGES_2026_27.xlsx', 'NMC_Colleges')
save_styled_excel(df_courses, 'data/NMC_COURSES_2026_27.xlsx', 'NMC_Courses')

print("\n--- Summary Verification ---")
print(f"Colleges File: data/NMC_COLLEGES_2026_27.xlsx -> {len(df_colleges)} rows")
print(f"Courses File:  data/NMC_COURSES_2026_27.xlsx  -> {len(df_courses)} rows")
print(f"Total Intake:  {df_courses['Annual_Intake'].sum():,} seats")
