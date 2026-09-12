"""
extract_coa_national.py — Council of Architecture (CoA) National Extractor
===========================================================================
Source: https://coa.gov.in/institutionStatus.php
Method: GET request, parses the official national directory of approved architectural institutions.
Extracts:
  - Official CoA Code (e.g. AP02, DL01, MH04, TS01)
  - Institution Name & Address
  - State & PIN Code
  - Affiliating University
  - Approved Courses (e.g. B.Arch, M.Arch) & Current Approved Intake
  - Aggregated to 1 physical institution record

Output:
  - data/raw/coa/coa_national_directory.html
  - data/raw/coa/coa_all_institutions.json
  - data/COA_INSTITUTIONS_2025.xlsx
"""

import sys, os
sys.path.insert(0, os.path.dirname(__file__))

import re
import json
from pathlib import Path
from datetime import date
from bs4 import BeautifulSoup

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE / "scripts"))

from lib_extract import (
    RateLimitedSession, load_checkpoint, save_checkpoint,
    RAW_DIR, RESEARCH_DIR, EXTRACTION_DATE
)

SOURCE = "coa"
RAW_SOURCE_DIR = RAW_DIR / SOURCE
RAW_SOURCE_DIR.mkdir(parents=True, exist_ok=True)

URL = "https://coa.gov.in/institutionStatus.php"

session = RateLimitedSession(
    rps=0.8,
    max_retries=5,
    timeout=30,
    extra_headers={
        "Accept": "text/html,application/xhtml+xml,*/*",
        "Referer": "https://coa.gov.in/",
    }
)

STATE_PREFIXES = {
    "AN": "Andaman and Nicobar Islands",
    "AP": "Andhra Pradesh",
    "AR": "Arunachal Pradesh",
    "AS": "Assam",
    "BR": "Bihar",
    "CH": "Chandigarh",
    "CG": "Chhattisgarh",
    "CT": "Chhattisgarh",
    "DL": "Delhi",
    "GA": "Goa",
    "GJ": "Gujarat",
    "HR": "Haryana",
    "HP": "Himachal Pradesh",
    "JK": "Jammu and Kashmir",
    "JH": "Jharkhand",
    "KA": "Karnataka",
    "KL": "Kerala",
    "LA": "Ladakh",
    "MP": "Madhya Pradesh",
    "MH": "Maharashtra",
    "MN": "Manipur",
    "ML": "Meghalaya",
    "MZ": "Mizoram",
    "NL": "Nagaland",
    "OD": "Odisha",
    "OR": "Odisha",
    "PY": "Puducherry",
    "PB": "Punjab",
    "RJ": "Rajasthan",
    "SK": "Sikkim",
    "TN": "Tamil Nadu",
    "TS": "Telangana",
    "TR": "Tripura",
    "UP": "Uttar Pradesh",
    "UA": "Uttarakhand",
    "UK": "Uttarakhand",
    "WB": "West Bengal",
}


def parse_coa_page(html_content: str) -> list[dict]:
    soup = BeautifulSoup(html_content, "html.parser")
    tables = soup.find_all("table")
    if not tables:
        return []

    table = tables[0]
    rows = table.find_all("tr")
    if len(rows) < 2:
        return []

    institutions_map = {}

    for row in rows[1:]:
        cells = [c.get_text(" ", strip=True) for c in row.find_all(["td", "th"])]
        if len(cells) < 4:
            continue

        raw_info = cells[1].strip()
        aff_univ = cells[2].strip() if len(cells) > 2 else ""
        course = cells[3].strip() if len(cells) > 3 else ""
        intake_info = cells[5].strip() if len(cells) > 5 else (cells[4].strip() if len(cells) > 4 else "")

        # Extract CoA Code at start of raw_info (e.g. AP02, DL01, TS03)
        code_m = re.match(r'^([A-Z]{2}\d{2,3})', raw_info)
        if not code_m:
            # Try searching anywhere in string
            code_m = re.search(r'\b([A-Z]{2}\d{2,3})\b', raw_info)
            
        coa_code = code_m.group(1) if code_m else ""
        if not coa_code:
            continue

        # Remainder after code
        rest = raw_info[len(coa_code):].strip() if raw_info.startswith(coa_code) else raw_info

        # Extract email, website, phone
        email_m = re.search(r'Email:\s*([a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+)', rest)
        email = email_m.group(1) if email_m else ""

        web_m = re.search(r'Website:\s*(https?://[^\s]+|[a-zA-Z0-9.-]+\.[a-zA-Z]{2,})', rest)
        website = web_m.group(1) if web_m else ""

        pin_m = re.search(r'\b(\d{6})\b', rest)
        pin_code = pin_m.group(1) if pin_m else ""

        # State from code prefix
        state_pfx = coa_code[:2]
        state = STATE_PREFIXES.get(state_pfx, "")

        # Clean inst name: usually between head designation and city
        # Or take the first meaningful segment
        lines = [l.strip() for l in rest.splitlines() if l.strip()]
        inst_name = lines[0] if lines else rest[:80]
        # Remove contact noise from inst_name
        inst_name = re.sub(r'(Tel:|Email:|Website:|Mobile:|Fax:).*', '', inst_name).strip()

        course_entry = f"{course} (Intake: {intake_info})" if intake_info else course

        if coa_code in institutions_map:
            if course_entry and course_entry not in institutions_map[coa_code]["courses"]:
                institutions_map[coa_code]["courses"].append(course_entry)
        else:
            institutions_map[coa_code] = {
                "official_institution_id": coa_code,
                "coa_code": coa_code,
                "institution_name": inst_name or f"Architectural Institute ({coa_code})",
                "full_text": rest,
                "state": state,
                "pin_code": pin_code,
                "university_affiliation": aff_univ,
                "email": email,
                "website": website,
                "country": "India",
                "education_level": "Higher Education / Architecture",
                "institution_type": "Architecture College / School of Planning & Architecture",
                "recognition_status": "Approved",
                "recognition_authority": "Council of Architecture (CoA)",
                "courses": [course_entry] if course_entry else [],
                "source_database": "Council of Architecture (CoA) Approved Institutions Register",
                "source_url": URL,
                "extraction_date": EXTRACTION_DATE,
            }

    return list(institutions_map.values())


def main():
    print(f"[CoA] Starting national extraction — {EXTRACTION_DATE}")
    status, raw = session.get(URL, as_json=False)
    if status != 200 or not raw:
        print(f"  Failed: status={status}")
        return

    html_content = raw.decode("utf-8", errors="replace") if isinstance(raw, bytes) else str(raw)
    
    # Save raw HTML
    raw_path = RAW_SOURCE_DIR / "coa_national_directory.html"
    raw_path.write_text(html_content, encoding="utf-8")
    print(f"  Saved raw HTML: {raw_path} ({len(html_content):,} chars)")

    records = parse_coa_page(html_content)
    print(f"  Extracted {len(records)} canonical architectural institutions nationally")

    # Save raw JSON
    json_path = RAW_SOURCE_DIR / "coa_all_institutions.json"
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=2)

    # Build Excel
    build_excel(records)


def build_excel(records: list[dict]):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        print("[WARN] openpyxl not available")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CoA_Institutions"

    headers = [
        "Official_ID", "CoA_Code", "Institution_Name",
        "State", "PIN_Code", "University_Affiliation",
        "Courses_and_Intake", "Email", "Website", "Country",
        "Education_Level", "Institution_Type",
        "Recognition_Status", "Recognition_Authority",
        "Academic_Year", "Source_URL", "Extraction_Date"
    ]

    hdr_fill = PatternFill("solid", fgColor="0B5345")
    hdr_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for i, r in enumerate(records, 2):
        row_data = [
            r.get("official_institution_id", ""),
            r.get("coa_code", ""),
            r.get("institution_name", ""),
            r.get("state", ""),
            r.get("pin_code", ""),
            r.get("university_affiliation", ""),
            "; ".join(r.get("courses", [])),
            r.get("email", ""),
            r.get("website", ""),
            r.get("country", "India"),
            r.get("education_level", "Higher Education / Architecture"),
            r.get("institution_type", "Architecture College"),
            r.get("recognition_status", "Approved"),
            r.get("recognition_authority", "Council of Architecture (CoA)"),
            "2025-26",
            r.get("source_url", URL),
            r.get("extraction_date", EXTRACTION_DATE),
        ]
        for col, v in enumerate(row_data, 1):
            ws.cell(row=i, column=col, value=v)

    col_widths = [15, 12, 45, 20, 10, 30, 40, 25, 30, 10, 25, 25, 15, 25, 12, 35, 15]
    for col, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w

    out_path = BASE / "data" / "COA_INSTITUTIONS_2025.xlsx"
    wb.save(out_path)
    print(f"  Excel saved: {out_path} ({len(records)} rows)")


if __name__ == "__main__":
    main()
