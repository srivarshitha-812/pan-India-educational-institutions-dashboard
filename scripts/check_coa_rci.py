import openpyxl
from pathlib import Path

def check_files():
    coa_p = Path("data/COA_INSTITUTIONS_2025.xlsx")
    wb_coa = openpyxl.load_workbook(coa_p, read_only=True)
    ws_coa = wb_coa.active
    coa_rows = list(ws_coa.iter_rows(values_only=True))
    wb_coa.close()
    
    headers_coa = coa_rows[0]
    data_coa = coa_rows[1:]
    h_map_coa = {h: i for i, h in enumerate(headers_coa)}
    id_idx_coa = h_map_coa.get("Official_ID", 0)
    coa_ids = [r[id_idx_coa] for r in data_coa if r[id_idx_coa]]
    
    print("=== COA VERIFICATION ===")
    print(f"File: {coa_p.name}")
    print(f"Total rows (excl header): {len(data_coa)}")
    print(f"Unique Official IDs: {len(set(coa_ids))}")
    print(f"Headers: {headers_coa}")
    
    rci_p = Path("data/RCI_INSTITUTIONS_2025.xlsx")
    wb_rci = openpyxl.load_workbook(rci_p, read_only=True)
    ws_rci = wb_rci.active
    rci_rows = list(ws_rci.iter_rows(values_only=True))
    wb_rci.close()
    
    headers_rci = rci_rows[0]
    data_rci = rci_rows[1:]
    h_map_rci = {h: i for i, h in enumerate(headers_rci)}
    id_idx_rci = h_map_rci.get("Official_ID", 0)
    rci_ids = [r[id_idx_rci] for r in data_rci if r[id_idx_rci]]
    
    print("\n=== RCI VERIFICATION ===")
    print(f"File: {rci_p.name}")
    print(f"Total rows (excl header): {len(data_rci)}")
    print(f"Unique Official IDs: {len(set(rci_ids))}")
    print(f"Headers: {headers_rci}")

if __name__ == "__main__":
    check_files()
